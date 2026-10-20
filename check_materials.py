"""Check that the investor materials agree with the operating model and with each other.

Run after changing Financial_Models/operating_model.py or editing any material:
    /opt/homebrew/bin/python3 check_materials.py
It reports every mismatch and exits non-zero if anything fails. Needs PyMuPDF (fitz) for PDFs;
.pptx text is read straight from the package XML, so python-pptx is not required.
"""
import re
import json
import sys
import zipfile
from html import unescape
from pathlib import Path

import fitz

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE / "Financial_Models"))
from operating_model import key_figures  # noqa: E402

IM = HERE / "Investor_Materials"
SEED_PPTX, SEED_PDF = IM / "Anvest_Seed_Deck.pptx", IM / "Anvest_Seed_Deck.pdf"
PARTNER_PPTX, PARTNER_PDF = IM / "Anvest_Strategic_Integration_Deck.pptx", IM / "Anvest_Strategic_Integration_Deck.pdf"
ONEPAGER, MEMO = IM / "Anvest_OnePager.pdf", IM / "Anvest_Data_Provenance_Memo.pdf"
EMAILS = HERE / "Outreach" / "Anvest_Outreach_Emails.md"
FUNDING, TARGETING = HERE / "Internal_Strategy" / "FUNDING_PLAN.md", HERE / "Internal_Strategy" / "INVESTOR_TARGETING_PLAN.md"

# Figures from earlier drafts that must not reappear in anything an investor sees.
STALE = ["₹8.3 L", "₹8.29", "₹51.8", "₹52 Cr", "₹50 Cr+", "₹14.5", "₹4.3 L", "₹4.3 lakh", "$5.4M", "$245M",
         "200 Desk seats", "900 paying", "2,100 research-analyst", "zero numeric hallucination", "●", "◐", "○",
         # Superseded claims and assumptions
         "₹83 L", "₹29.9", "$247M", "₹95.69", "$330K", "investing co-pilot", "9 in 10", "F&O traders lost",
         "100% exact", "100% match", "100% source", "outreach underway", "pilots underway", "unmetered",
         "ILLUSTRATIVE", "Illustrative example", "~17K", "~17,000", "17,300", "15,718", "~100", "5,405",
         "6,137 stock pages with AI profiles", "2.89M", "Only Anvest", "zero-CAC", "Zero-CAC", "Product capability (live today)", "₹13 Cr"]
# Speaker notes travel with every .pptx: nothing internal may sit in them.
NOTES_FORBIDDEN = re.compile(r"\.py\b|\.md\b|Internal_Strategy|Financial_Models|VERIFY|[Vv]erify before|TODO|"
                             r"data check pending|returned an error|bug\b", re.I)
VERSION_RX = re.compile(r"20 Oct 2026")


def squash(t):
    return re.sub(r"\s+", " ", t)


def pptx_text(path):
    with zipfile.ZipFile(path) as z:
        parts = [n for n in z.namelist() if re.match(r"ppt/(slides|slideLayouts)/[^/]+\.xml$", n)]
        return squash(" ".join(unescape(t) for n in parts for t in re.findall(r"<a:t>([^<]*)</a:t>", z.read(n).decode("utf8"))))


def notes_text(path):
    with zipfile.ZipFile(path) as z:
        parts = [n for n in z.namelist() if re.match(r"ppt/notesSlides/[^/]+\.xml$", n)]
        return squash(" ".join(unescape(t) for n in parts for t in re.findall(r"<a:t>([^<]*)</a:t>", z.read(n).decode("utf8"))))


def pdf_text(path):
    return squash(" ".join(p.get_text() for p in fitz.open(path)))


def inr(cr):
    """Same rupee format as the decks: lakh below ₹1 Cr, two decimals (trailing zeros dropped) below ₹10 Cr."""
    if cr < 1:
        return f"₹{round(cr * 100)} L"
    return "₹" + (f"{cr:.0f}" if cr >= 10 else f"{cr:.2f}".rstrip("0").rstrip(".")) + " Cr"


def check_outreach(k, failures):
    """Reconcile the contact register, CRM export and round terms; no network writes."""
    from openpyxl import load_workbook
    register = json.loads((HERE / "Outreach/investor_contacts.json").read_text())
    records = register["targets"]
    if len({r["id"] for r in records}) != len(records):
        failures.append("Outreach contact register has duplicate record IDs")
    book = load_workbook(HERE / "Outreach/Anvest_Investor_Tracker.xlsx", data_only=True)
    if book.sheetnames != ["DO THIS NEXT", "Metrics Gate", "Investor Targets", "Round Model", "Timeline", "Sources"]:
        failures.append("Outreach tracker changed its six-tab workflow")
    targets = book["Investor Targets"]
    actual = {row[28]: row for row in targets.iter_rows(min_row=6, max_col=29, values_only=True) if row[28]}
    if set(actual) != {r["id"] for r in records}:
        failures.append("Outreach workbook and contact register contain different target IDs")
    for r in records:
        row = actual.get(r["id"])
        if not row:
            continue
        for i, field in ((1, "name"), (16, "email"), (18, "url"), (19, "source")):
            if (row[i] or "") != r[field]:
                failures.append(f"Outreach {r['id']}: exported {field} differs from contact register")
        if r["email"] and not r["source"].startswith("https://"):
            failures.append(f"Outreach {r['id']}: published email lacks a primary source")
        if row[24] is not None and (not isinstance(row[24], (int, float)) or not 0 <= row[24] <= 2):
            failures.append(f"Outreach {r['id']}: follow-up count must be 0, 1 or 2")
    round_sheet = book["Round Model"]
    for c, case in (("B", "target"), ("C", "first_close"), ("D", "maximum")):
        for row, expected in ((6, k[case]["raise_usd_k"] * 1000), (8, k["fx"]),
                              (25, k[case]["arr_cr"] * 1e7), (26, k[case]["teams"]),
                              (27, k[case]["seats"]), (28, k[case]["partners"]),
                              (29, k[case]["breakeven_month"])):
            if round_sheet[f"{c}{row}"].value != expected:
                failures.append(f"Outreach round {c}{row}: does not match operating model")
    if round_sheet["B22"].value != 0 or round_sheet["B21"].value != k["target"]["raise_usd_k"] * 1000:
        failures.append("Outreach target use-of-funds allocation does not reconcile")
    usage = json.loads((HERE / "Source/product_metrics_snapshot.json").read_text())["users_and_usage"]
    metrics = {row[0]: row[3] for row in book["Metrics Gate"].iter_rows(min_row=6, max_col=6, values_only=True) if row[0]}
    for label, expected in (("Registered accounts", usage["registered"]),
                            ("External accounts with workflow", usage["activated_external"]),
                            ("External activity in 2+ weeks", usage["retained_external_2plus_weeks"])):
        if metrics.get(label) != expected:
            failures.append(f"Outreach metric {label}: differs from dated product snapshot")
    for sheet in book:
        for row in sheet:
            for cell in row:
                if cell.data_type == "e":
                    failures.append(f"Outreach {sheet.title}!{cell.coordinate}: {cell.value}")
    book.close()


def main():
    k = key_figures()
    f, t, x, mk = k["first_close"], k["target"], k["maximum"], k["market"]
    expected = {
        "$500K month-15 ARR": inr(t["arr_cr"]),
        "$400K month-15 ARR": inr(f["arr_cr"]),
        "$700K month-15 ARR": inr(x["arr_cr"]),
        "$500K Desk teams": str(t["teams"]),
        "$500K break-even": f"M{t['breakeven_month']}",
        "TAM": f"${mk['tam_usd_b']}B",
        "SAM": f"${mk['sam_usd_m']}M",
    }
    failures = []
    for pdf in (SEED_PDF, PARTNER_PDF, ONEPAGER, MEMO):
        if not pdf.exists():
            failures.append(f"{pdf.name} is missing: export it before sending")
    if failures:
        print("FAIL"); [print("  -", m) for m in failures]; sys.exit(1)
    texts = {
        "seed deck .pptx": pptx_text(SEED_PPTX), "seed deck .pdf": pdf_text(SEED_PDF),
        "partner deck .pptx": pptx_text(PARTNER_PPTX), "partner deck .pdf": pdf_text(PARTNER_PDF),
        "one-pager": pdf_text(ONEPAGER), "memo": pdf_text(MEMO),
        "outreach emails": squash(EMAILS.read_text()),
        "FUNDING_PLAN.md": squash(FUNDING.read_text()), "INVESTOR_TARGETING_PLAN.md": squash(TARGETING.read_text()),
    }
    needs = {
        "seed deck .pptx": list(expected), "seed deck .pdf": list(expected),
        "one-pager": ["$500K month-15 ARR", "$700K month-15 ARR", "$500K Desk teams", "$500K break-even", "TAM", "SAM"],
        "FUNDING_PLAN.md": ["$500K month-15 ARR", "$400K month-15 ARR", "$700K month-15 ARR", "$500K Desk teams"],
        "INVESTOR_TARGETING_PLAN.md": ["$500K month-15 ARR", "$400K month-15 ARR", "$700K month-15 ARR", "SAM"],
    }
    for doc, keys in needs.items():
        for key in keys:
            if expected[key] not in texts[doc]:
                failures.append(f"{doc}: missing {key} ({expected[key]})")
    for doc in ("seed deck .pptx", "seed deck .pdf", "partner deck .pptx", "partner deck .pdf", "one-pager", "memo", "outreach emails"):
        for s in STALE:
            if s in texts[doc]:
                failures.append(f"{doc}: stale or unsupported text '{s}'")
    for doc, pptx in (("seed deck notes", SEED_PPTX), ("partner deck notes", PARTNER_PPTX)):
        for m in sorted(set(NOTES_FORBIDDEN.findall(notes_text(pptx)))):
            failures.append(f"{doc}: internal or unsupported text '{m}'")
    # Dated production evidence and source-audit scope must survive every export.
    for doc in ("seed deck .pptx", "seed deck .pdf", "one-pager"):
        for claim in ("50 / 50", "418", "238", "31"):
            if claim not in texts[doc]:
                failures.append(f"{doc}: missing scoped traction evidence {claim!r}")
        if not any(scope in texts[doc] for scope in ("founder/testing", "including internal testing")):
            failures.append(f"{doc}: usage totals omit the internal-testing qualification")
    for doc in ("seed deck .pptx", "seed deck .pdf", "partner deck .pptx", "partner deck .pdf", "one-pager", "memo"):
        if re.search(r"\br[1-9]\d*\b", texts[doc]):
            failures.append(f"{doc}: internal draft label leaked")
    for doc in ("seed deck .pptx", "seed deck .pdf"):
        for claim in ("₹13.4 Cr", "₹9.7 Cr", "₹23.1 Cr", "team workspaces planned", "19%", "CA/wealth", "67 bps", "total income"):
            if claim not in texts[doc]: failures.append(f"{doc}: missing reconciliation {claim!r}")
    audit=json.loads((HERE/"Source/filing_check.json").read_text())
    if audit.get("passed") != 50 or audit.get("total") != 50 or any(r.get("status") != "pass" for r in audit.get("results",[])):
        failures.append("The 50-filing audit is incomplete or failed")
    if len(audit.get("results",[])) != 50 or len({r["symbol"] for r in audit.get("results",[])}) != 50:
        failures.append("The filing audit does not contain 50 distinct companies")
    check_outreach(k, failures)
    # Compare the exported snapshot against the Python schedule (not just labels).
    from openpyxl import load_workbook
    from operating_model import SCENARIO, run
    wb=load_workbook(HERE/"Financial_Models/Anvest_Operating_Model.xlsx",read_only=True,data_only=True)
    ws=wb["Scenario"]
    for r in run(SCENARIO):
        for col,expected_value in ((14,round(r.cost["salaries"])),(17,round(r.cost["growth"])),(19,round(r.spend)),(20,round(r.cash))):
            if ws.cell(r.m+1,col).value != expected_value:
                failures.append(f"Workbook scenario month {r.m}: stale cost/cash in column {col}")
    wb.close()
    versions = {doc: set(VERSION_RX.findall(texts[doc])) for doc in
                ("seed deck .pptx", "seed deck .pdf", "partner deck .pptx", "partner deck .pdf", "one-pager", "memo")}
    labels = set().union(*versions.values())
    if len(labels) != 1 or any(len(v) != 1 for v in versions.values()):
        failures.append(f"version labels disagree or are missing: {versions}")
    for pptx, pdf in ((SEED_PPTX, SEED_PDF), (PARTNER_PPTX, PARTNER_PDF)):
        if pdf.stat().st_mtime + 1 < pptx.stat().st_mtime:
            failures.append(f"{pdf.name} is older than {pptx.name}: re-export the PDF")
    label = next(iter(labels)) if len(labels) == 1 else "?"
    print(f"Materials version: {label}")
    print("Model month-15 ARR: $400K " + expected["$400K month-15 ARR"] + ", $500K " + expected["$500K month-15 ARR"]
          + ", $700K " + expected["$700K month-15 ARR"] + f", $500K break-even {expected['$500K break-even']}"
          + f", year-5 scenario {inr(k['scenario_arr_cr']['Y5'])}, TAM {expected['TAM']}, SAM {expected['SAM']}")
    if failures:
        print(f"FAIL ({len(failures)})")
        for f in failures:
            print("  -", f)
        sys.exit(1)
    print("PASS: decks, PDFs, one-pager, memo, emails and plans agree with the operating model, carry one version label, "
          "and the speaker notes hold nothing internal. Source-audit scope and scenario workbook costs match.")


if __name__ == "__main__":
    if "--outreach" in sys.argv:
        issues = []
        check_outreach(key_figures(), issues)
        if issues:
            print("FAIL: outreach consistency")
            for issue in issues:
                print("  -", issue)
            sys.exit(1)
        print("PASS: outreach contacts, source IDs, six-tab workflow, formulas, dated metrics and round terms reconcile.")
    else:
        main()
